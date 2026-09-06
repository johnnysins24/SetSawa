import io
import json
import subprocess
import sys
from pathlib import Path

import pytest

from setsawa.core import clean_cells, contained, csv_value, normalize, scoped_id
from setsawa.sandbox import parse_pdf


def cells():
    return ["1", "123xxxxxx4567", "ชื่อ ทดสอบ", "+12", "อบต.ทดสอบ", "ข้อความต้นทาง", "อำเภอหนึ่ง", "จังหวัดหนึ่ง", "ข้อความยืนยันต้นทาง", "สถานะต้นทาง"]


def test_duplicate_rule_ignores_only_ordinal():
    first, second = cells(), cells()
    second[0] = "2"
    assert clean_cells(first) == clean_cells(second)
    for index in range(1, 10):
        altered = cells()
        altered[index] = "124xxxxxx4567" if index == 1 else "+13" if index == 3 else altered[index] + "อื่น"
        assert clean_cells(altered)["record_id"] != clean_cells(first)["record_id"]


def test_mask_is_not_identity_key():
    first, second = cells(), cells()
    second[2] = "อีกคน ชื่ออื่น"
    assert clean_cells(first)["masked_id"] == clean_cells(second)["masked_id"]
    assert clean_cells(first)["record_id"] != clean_cells(second)["record_id"]


def test_geography_is_scoped():
    assert scoped_id("district", "ก", "เมือง") != scoped_id("district", "ข", "เมือง")
    assert scoped_id("agency", "ก", "อ", "อบต.เดียว") != scoped_id("agency", "ก", "บ", "อบต.เดียว")


def test_thai_and_unicode():
    assert normalize("  ชื่อ\n\tนามสกุล  ") == "ชื่อ นามสกุล"
    assert normalize("é") == normalize("e\u0301")
    assert normalize("ตําบล") == "ตําบล"  # no blanket NFKC or spelling corrections


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-cmd", "@SUM(A1)", "\t=1", "\r@cmd", "\ufeff=1"])
def test_csv_formula_injection_rejected(value):
    with pytest.raises(ValueError):
        csv_value(value)
    assert csv_value(12) == 12


@pytest.mark.parametrize("path", ["../escape", "a/../../escape", "/absolute", "C:/escape", "data:stream"])
def test_path_traversal_rejected(tmp_path, path):
    with pytest.raises(ValueError):
        contained(tmp_path, path)


def test_malformed_pdf_rejected(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF-1.7\nnot a pdf")
    with pytest.raises(ValueError, match="PDF rejected"):
        parse_pdf(path)


def test_active_pdf_rejected(tmp_path):
    from pypdf import PdfWriter
    path = tmp_path / "active.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_js("app.alert('test')")
    with path.open("wb") as stream:
        writer.write(stream)
    with pytest.raises(ValueError, match="Active content"):
        parse_pdf(path)


def test_pdf_worker_timeout():
    from setsawa.core import ROOT
    path = next((ROOT / "data/raw/pdfs").glob("01_*.pdf"))
    with pytest.raises(ValueError, match="time/output limit"):
        parse_pdf(path, timeout=0.001)


def test_network_and_subprocess_denied():
    from setsawa.pdf_worker import restricted
    for event in ("socket.__new__", "socket.connect", "subprocess.Popen", "os.system"):
        with pytest.raises(PermissionError):
            restricted(event, ())


def test_xlsx_literal_injection(tmp_path):
    import xlsxwriter
    from openpyxl import load_workbook
    from setsawa.workbook import put_literal
    path = tmp_path / "literal.xlsx"
    with xlsxwriter.Workbook(path) as book:
        sheet = book.add_worksheet()
        put_literal(sheet, 0, 0, '=HYPERLINK("https://example.invalid","test")')
    book = load_workbook(path)
    assert book.active["A1"].data_type == "s"
    assert book.active["A1"].value.startswith("=HYPERLINK")
    book.close()


def test_secret_scan(tmp_path):
    from setsawa.security import scan_files
    path = tmp_path / "test.txt"
    path.write_text("ghp_" + "A" * 36)
    with pytest.raises(ValueError, match="Potential secret"):
        scan_files([path])


def test_missing_cells_and_invalid_mask_rejected():
    for i, value in ((1, "1234567890123"), (2, ""), (3, "NaN"), (4, None)):
        row = cells()
        row[i] = value
        with pytest.raises(ValueError):
            clean_cells(row)
