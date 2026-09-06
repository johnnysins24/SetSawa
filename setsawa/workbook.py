"""Standard-Python workbook export for the user-requested reproducible CLI."""
import datetime as dt
import math
from pathlib import Path

from .core import FIELDS, HEADERS_TH, NOTICE, NOTICE_TH, read_json

SHEETS = {"records": "Records", "provenance": "Provenance", "duplicates": "Duplicates", "provinces": "Provinces", "districts": "Districts", "agency_locations": "Agency locations"}
LABELS = dict(zip(FIELDS, HEADERS_TH, strict=True)) | {
    "source_status_text": "สถานะตามต้นทาง", "source_verification_status_text": "สถานะการยืนยันตามต้นทาง",
    "record_id": "รหัสระเบียน", "source_ordinal": "ลำดับต้นทาง", "page_number": "หน้า", "data_row_on_page": "แถวข้อมูลในหน้า",
    "record_count": "ระเบียนไม่ซ้ำ", "source_row_count": "แถวต้นทาง", "record_score_sum": "รวมคะแนนหลังตัดซ้ำ", "source_score_sum": "รวมคะแนนต้นทาง",
}


def put_literal(sheet, row, col, value, fmt=None):
    if isinstance(value, str):
        sheet.write_string(row, col, value, fmt)
    elif isinstance(value, (int, float)):
        sheet.write_number(row, col, value, fmt)
    else:
        raise TypeError("Unsupported workbook cell")


def make_workbook(output):
    import xlsxwriter
    target = output / "reader/setsawa.xlsx"
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(target, {"strings_to_formulas": False, "strings_to_urls": False})
    workbook.set_properties({"title": "SetSawa 2568 / 2025", "subject": "Provided source documents and cleaned transcription", "author": "SetSawa", "created": dt.datetime(2026, 9, 6)})
    title = workbook.add_format({"font_name": "Arial", "font_size": 17, "bold": True, "font_color": "#172E48"})
    body = workbook.add_format({"font_name": "Arial", "font_size": 11, "valign": "top"})
    note = workbook.add_format({"font_name": "Arial", "font_size": 11, "text_wrap": True, "valign": "top", "font_color": "#475569"})
    integer = workbook.add_format({"font_name": "Arial", "font_size": 11, "num_format": "#,##0", "valign": "top"})
    header_format = workbook.add_format({"font_name": "Arial", "font_size": 10, "bold": True, "font_color": "#FFFFFF", "bg_color": "#203B5C", "text_wrap": True, "valign": "vcenter"})
    start = workbook.add_worksheet("Read me")
    start.hide_gridlines(2)
    start.set_column("A:A", 27)
    start.set_column("B:B", 105, note)
    start.write_string(0, 0, "SetSawa", title)
    content = [("Collection year", "2568 BE / 2025 CE; PDF creation and processing dates are separate."), ("Source rows / แถวต้นทาง", "3,966 rows in 77 supplied PDFs, 149 pages."), ("Distinct records / ระเบียน", "3,963 after three exact content duplicate collapses."), ("Meaning", NOTICE), ("ข้อจำกัด", NOTICE_TH), ("How to read", "Records contains the clean snapshot. Join record_id to Provenance to find every PDF page and original row. Duplicates documents all collapses."), ("Excel headers", "Each table has Thai descriptions on row 3 and stable English field names on row 4."), ("IDs", "Existing ID masks remain unchanged. record_id identifies a content record, not an independently verified person."), ("Geography", "District and agency groups include province and district labels. They are not official administrative codes."), ("Source data reuse", "Unspecified. The MIT code license does not license the source documents or their contents."), ("Repository", "https://github.com/johnnysins24/SetSawa"), ("Source collection", "รวมมิตรท้องถิ่น 68; see Sources for byte hashes and preserved relative file paths."), ("Snapshot", "1.0.0; summaries are precomputed from the released data. Rebuild with the CLI after edits.")]
    for n, (label, value) in enumerate(content, 3):
        start.write_string(n, 0, label, body)
        start.write_string(n, 1, value, note)
        start.set_row(n, 44 if len(value) > 100 else 30)

    def add_table(sheet_name, rows, table_name):
        sheet = workbook.add_worksheet(sheet_name)
        sheet.hide_gridlines(2)
        sheet.write_string(0, 0, sheet_name, title)
        sheet.set_row(0, 28)
        sheet.set_row(2, 44)
        sheet.set_row(3, 44)
        keys = list(rows[0])
        if sheet_name == "Records":
            front = ["full_name_th", "province_name_th", "district_name_th", "agency_name_th", "score_adjustment_reported", "masked_id"]
            keys = front + [k for k in keys if k not in front]
        for c, key in enumerate(keys):
            sheet.write_string(2, c, LABELS.get(key, key.replace("_", " ")), note)
            width = 24 if "id" in key else 32
            if "name_th" in key:
                width = 32
            if key in ("source_location_basis_text", "pdf_relative_path", "relative_path"):
                width = 62
            if isinstance(rows[0][key], (int, float)):
                width = 25
            if key in {"record_id", "province_id", "district_id", "agency_location_id", "pdf_sha256", "sha256"}:
                width = 78
            width = max(width, min(58, len(key) + 4))
            sheet.set_column(c, c, width, integer if isinstance(rows[0][key], int) else body)
        sheet.add_table(3, 0, len(rows) + 3, len(keys) - 1, {"name": table_name, "style": "Table Style Medium 2", "columns": [{"header": k, "header_format": header_format} for k in keys]})
        for ri, row in enumerate(rows, 4):
            for ci, key in enumerate(keys):
                put_literal(sheet, ri, ci, row[key], integer if isinstance(row[key], int) else body)
        sheet.freeze_panes(4, 0)
        sheet.set_landscape()
        sheet.fit_to_pages(1, 0)
        sheet.repeat_rows(2, 3)
        return sheet

    for table, name in SHEETS.items():
        add_table(name, read_json(output / f"clean/{table}.json"), table.replace("_", "") + "Data")
    from .core import ROOT
    schema = read_json(ROOT / "schemas/records.schema.json")
    dictionary = [{"field": k, "thai_label": LABELS.get(k, k), "type": v["type"], "description": v["description"]} for k, v in schema["properties"].items()]
    ds = add_table("Dictionary", dictionary, "DictionaryData")
    ds.set_column("D:D", 88, note)
    for row in range(4, len(dictionary) + 4):
        ds.set_row(row, 44)
        ds.write_string(row, 3, dictionary[row - 4]["description"], note)
    sources = [{"document_id": d["document_id"], "filename": d["filename"], "relative_path": d["relative_path"], "sha256": d["sha256"], "page_count": d["page_count"], "source_row_count": d["summary"]["source_row_count"]} for d in read_json(output / "raw/documents.json")]
    add_table("Sources", sources, "SourcesData")
    workbook.close()
    print("Wrote reader workbook with nine sheets")


def verify_workbook(path, tables):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if workbook.sheetnames != ["Read me", *SHEETS.values(), "Dictionary", "Sources"]:
            raise ValueError("Unexpected workbook sheet set")
        for table, name in SHEETS.items():
            sheet = workbook[name]
            rows = tables[table]
            header = [c.value for c in sheet[4]]
            actual = [dict(zip(header, row, strict=True)) for row in sheet.iter_rows(min_row=5, values_only=True)]
            equivalent = len(actual) == len(rows) and len(header) == len(rows[0]) and set(header) == set(rows[0])
            if equivalent:
                for a, b in zip(actual, rows, strict=True):
                    for key in header:
                        if key.startswith("bbox_"):
                            equivalent = equivalent and isinstance(a[key], (float, int)) and math.isclose(a[key], b[key], rel_tol=0, abs_tol=1e-9)
                        else:
                            equivalent = equivalent and a[key] == b[key]
            if not equivalent or sheet.freeze_panes != "A5" or not sheet.tables:
                raise ValueError("Workbook data/layout mismatch: " + name)
        for sheet in workbook:
            for row in sheet:
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Unexpected formula in literal snapshot workbook")
    finally:
        workbook.close()
